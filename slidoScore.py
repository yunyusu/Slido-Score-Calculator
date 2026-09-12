#! python3
# slidoScore.py - calculate the quiz score for syudents in 腦與心智 course
# input: 對Slido匯出的Poll results per participant的Excel檔按滑鼠右鍵，
#        選「複製路徑」後回到程式頁面，按「F5」執行檔案
# output: (1) 列出未正確輸入個資的同學
#         (2) 列出有部分題目錯過的同學
#         (3) 將當週小考成績寫入成績計算Excel表(寫入時不可打開該Excel表)
# Author: 蘇勻郁
# Last update: 2026/9/12
# Updating content: 
# Wait for update: 移到colab或者做圖形化介面

import pyperclip, re, openpyxl, sys, os
from openpyxl.styles import Font

## parameter setting ###################################################################################
quiz_num = 1       # 這次是第幾次quiz
quiz_total_num = 15 # 這學期總共有幾次quiz
quiz_count = 12     # 這學期quiz取幾次最高分計算成績

# 小考成績登記表在電腦中的路徑
scoreTable_addr = "C://Users//yunyu//Desktop//大學//助教工作//腦與心智//小考成績登記.xlsx" 
sheet_name = "工作表1" # 工作表名稱
first_time_setting = 1 # 第一次跑此程式，須設定表格第一行標題: 1(開啟)/0關閉
deleteAll = 1          # 打掉重來，刪掉所有表格現有資料: 1(開啟)/0關閉

# 第一次上課可能還沒有統計好Email，需要關掉部分身分偵測
name_no_mercy = 1      # 名字打錯偵測(打錯=零分): 1(開啟)/0關閉
id_no_mercy = 1        # ID打錯偵測(打錯=零分): 1(開啟)/0關閉
email_no_mercy = 1     # Email打錯偵測(打錯=零分): 1(開啟)/0關閉

# 送分題設定(漏答也加分，但該堂課未到/個資不完整者不加分): 1(開啟)/0關閉
giveAway = [0, 0, 0, 0, 0] # [Q1, Q2, Q3, Q4, Q5]
########################################################################################################

addr = pyperclip.paste()
addr = addr[1: len(addr)-1]

addr_sep = addr.split(os.path.sep)
filename = addr_sep[len(addr_sep)-1]
print("Calculating score for the excel file: " + filename + " ......\n")

wb = openpyxl.load_workbook(addr)
sheet = wb['Pivot All']
student_score = []

# Create a new student list that excludes those without basic information
# 1. ID Test: 是否為清大學號、是否有中英文字母、字數是否符合
idNumRegex_front = re.compile(r'^1\d\d\d\d\d\d\d\d')
idNumRegex_back = re.compile(r'1\d\d\d\d\d\d\d\d$')
# 2. Name Test: 是否含學號、空白
nameRegex_front = re.compile(r'^(\D)+')
nameRegex_back = re.compile(r'(\D)+$')
# 3. Email Test: 是否空白、符合Email格式
emailRegex = re.compile(r'(\w)+@(\w)+')

student_infoLoss = []
student_missQuestion = []
for rowNum in range(3, sheet.max_row + 1):
    id_test = False
    name_test = False
    email_test = False
    # print(sheet.cell(row=rowNum, column=1).value)
    # 1. ID Test
    if (sheet.cell(row=rowNum, column=4).value != None and sheet.cell(row=rowNum, column=4).value != ""):
        mo1 = idNumRegex_front.search(sheet.cell(row=rowNum, column=4).value)
        mo2 = idNumRegex_back.search(sheet.cell(row=rowNum, column=4).value)
        id_test = (mo1 == None or mo2 == None)
    else: id_test = True
    # 2. Name Test
    if (sheet.cell(row=rowNum, column=2).value != None and sheet.cell(row=rowNum, column=2).value != ""):
        mo1 = nameRegex_front.search(sheet.cell(row=rowNum, column=2).value)
        mo2 = nameRegex_back.search(sheet.cell(row=rowNum, column=2).value)
        name_test = (mo1 == None or mo2 == None)
    else: name_test = True
    # 3. Email Test
    if (sheet.cell(row=rowNum, column=3).value != None and sheet.cell(row=rowNum, column=3).value != ""):
        mo1 = emailRegex.search(sheet.cell(row=rowNum, column=3).value)
        email_test = (mo1 == None)
    else: email_test = True
    # Compare all the test
    if ((name_no_mercy and name_test) or (id_no_mercy and id_test) or (email_no_mercy and email_test)):
        # Somethong wrong with the identity information
        #error_code = 1*id_test + 2*name_test + 4*email_test
        # 0:all correct, 1:id, 2:name, 3:id & name. 4:email, 5:id & email, 6:name & email 7: id & name & email
        #match error_code:
            #case 1:
                #print("[" +  + "]")
        error_message = "["
        if (name_test):
            if (sheet.cell(row=rowNum, column=2).value != None and sheet.cell(row=rowNum, column=2).value != ""):
                error_message = error_message + sheet.cell(row=rowNum, column=2).value
            else: error_message = error_message + "???"
        else:
            error_message = error_message + sheet.cell(row=rowNum, column=2).value
        error_message = error_message + ", "
        if (id_test):
            error_message = error_message + "???"
        else:
            error_message = error_message + sheet.cell(row=rowNum, column=4).value
        error_message = error_message + ", "
        if (email_test):
            error_message = error_message + "???"
        else:
            error_message = error_message + sheet.cell(row=rowNum, column=3).value
        error_message = error_message + "]"
        # print(error_message)
        student_infoLoss.append(error_message)
        # 若想刪掉個資不完全的資料，可取消下方註解
        # sheet.delete_rows(rowNum, 1)
    else:
        # The information is complete and perfect!
        # row_value = ['姓名', '學號', '信箱', '分數', '漏答題號']
        row_value = []
        row_value.append(sheet.cell(row=rowNum, column=2).value) # name
        row_value.append(sheet.cell(row=rowNum, column=4).value) # id
        row_value.append(sheet.cell(row=rowNum, column=3).value) # email
        # check the answering history for Q1~Q5
        question_miss = []
        for columnNum in range(6, 11):
            if (sheet.cell(row=rowNum, column=columnNum).value == ""):
                if (giveAway[columnNum-6] == 0):
                    question_miss.append(str(columnNum-5))
                else:
                    # 漏答者也可加分
                    sheet.cell(row=rowNum, column=5).value = int(sheet.cell(row=rowNum, column=5).value) + 1
            elif (giveAway[columnNum-6] == 1):
                # 判斷同學原本是否就有答對，沒有加分，有就維持不變
                text = sheet.cell(row=rowNum, column=columnNum).value
                text = text.split(" ")
                if (not text[len(text)-1] == '(correct)'):
                    sheet.cell(row=rowNum, column=5).value = int(sheet.cell(row=rowNum, column=5).value) + 1
            else: continue
        row_value.append(int(sheet.cell(row=rowNum, column=5).value)) # score
        miss_list = ", ".join(question_miss)
        if (len(question_miss) != 0):
            # print(miss_list)
            row_value.append(miss_list)
            student_missQuestion.append("[" + sheet.cell(row=rowNum, column=2).value + ", "
                                        + sheet.cell(row=rowNum, column=4).value + ", "
                                        + sheet.cell(row=rowNum, column=3).value + "]: "
                                        + miss_list)
        # print(row_value)
        student_score.append(row_value)

# Find students without student ID, name , or email => 直接零分，解除註解可看名單提供預警
print("*************************************************************")
print("There are " + str(len(student_infoLoss)) + " students without full idetity:")
for i in range(len(student_infoLoss)):
    print(student_infoLoss[i])

# Find students that miss questions (no response) => 統一回覆：15次取12次下次再加油！可考慮發脆加分
print("*************************************************************")
print("There are " + str(len(student_missQuestion)) + " students miss answering questions:")
for i in range(len(student_missQuestion)):
    print(student_missQuestion[i])

# The information of students with complete identity information => 解除註解可看丟進excel前的資料長相
print("*************************************************************")
print("There are " + str(len(student_score)) + " students score recorded...")
# for i in range(len(student_score)):
#     print(student_score[i])

# Score Table
# 姓名   | 學號      | 信箱          | 最終成績 | Quiz 1 | Quiz 2 | ... | Quiz 15 |
# 窩豪幫 | 112000000 | abc@gmail.com | 100      | 5      | 5      | ... | 5       |
# ...

wb_score = openpyxl.load_workbook(scoreTable_addr)
sheet_score = wb_score[sheet_name]

# 刪去整個表格，打掉重練
if (deleteAll):
    sheet_score.delete_rows(1, sheet_score.max_row)

# 建立小考成績登記Excel表第一行標題(如上格式)
header = ["姓名", "學號", "信箱", "最終成績"]
if (first_time_setting):
    for columnNum in range(1, len(header)+1):
        sheet_score.cell(1, columnNum).value = header[columnNum-1]
    for i in range(1, quiz_total_num+1):
        sheet_score.cell(1, i+4).value = "Quiz " + str(i)

# quiz_num有無更新的防呆
# 若無資料 => 繼續程序
# 已有資料 => 再次確認，若不對則更新quiz_num，對則刪去原有資料再繼續程序
columnUsed = True
value_check = 0
while(columnUsed):
    for rowNum in range(2, sheet_score.max_row + 1):
        # print(sheet.cell(rowNum, quiz_num+4).value)
        value_check = value_check + sheet_score.cell(rowNum, quiz_num+4).value
    if (value_check == 0): columnUsed = False
    else:
        while (value_check != 0):
            print("You're now recording Quiz " + str(quiz_num) + " score, do you want to change quiz_num?")
            temp = input("Recheck your quiz_num, type again: ")
            # The score for this quiz will be re-evaluated
            if (quiz_num == int(temp)):
                quiz_num = int(temp)
                print("Quiz " + str(quiz_num) + " old record will be erased and covered with new data.")
                for rowNum in range(2, sheet_score.max_row + 1):
                    sheet_score.cell(rowNum, quiz_num+4).value = 0
                value_check = 0
                columnUsed = False
            # assign new quiz_num, recheck if it's vacant
            elif (int(temp) > 0 and int(temp) < quiz_total_num):
                quiz_num = int(temp)
                value_check = 0
            # the input number is out of range, reassigned one
            else:
                print("The number is out of range.")

addr_sep = scoreTable_addr.split(os.path.sep)
filename = addr_sep[len(addr_sep)-1]
print("Quiz " + str(quiz_num) + " score is recorded to file: " + filename + "......")

# If the name exist, update the new quiz score to the row; else, append to the last row
# row_value = ['姓名', '學號', '信箱', '分數', '漏答題號']
# Go through student list, find matching student ID list
student_for_check = []
for i in range(len(student_score)):
    findStudent = False
    max_row = sheet_score.max_row + 1
    for rowNum in range(2, max_row):
        # 比對學號
        if (sheet_score.cell(row=rowNum, column=2).value == student_score[i][1]):
            # find existing student
            # (1) student may repeatedly log in and out and in within the class
            # Then, his/her score is now recorded as the first log in record
            if (sheet_score.cell(rowNum, quiz_num+4).value != 0):
                # check if his/her is already on the check list
                for k in range(len(student_for_check)):
                    if (student_for_check[k] == student_score[i][1]):
                        findStudent = True
                if (not findStudent):
                    student_for_check.append(student_score[i][1])
                    # print(f"Find existing student: {student_score[i][1]}, check individually for his/her score.")
                findStudent = True
                break
            # (2) student is recorded in class in other quizzes, but haven't have the score for the latest quiz
            else:
                sheet_score.cell(rowNum, quiz_num+4).value = student_score[i][3]
                # print(f"Find existing student: {student_score[i][1]}, the score for the new quiz is recorded.")
    # cannot find in student score list, append to the last row
    # 若已於登記表上填入修課同學名單且只想登記有修課同學的成績，可將下方改為註解
    if (not findStudent):
        sheet_score.cell(max_row, 1).value = student_score[i][0]
        sheet_score.cell(max_row, 2).value = student_score[i][1]
        sheet_score.cell(max_row, 3).value = student_score[i][2]
        sheet_score.cell(max_row, quiz_num+4).value = student_score[i][3]

# check the original student data from slido sheet for students with problem 
# #1: 112123456
#                1 0 1 0 0
#                0 1 0 0 0
# #2: 112345678
# ......
# 1 = correct, 0 = incorrect
# print(student_for_check)
print("*************************************************************")
print("Check individually for students' who have multiple record in a single quiz.\n This can be cause by accidentally log in and out within class.")
for i in range(len(student_for_check)):
    print(f"#{i+1:02d}: {student_for_check[i]}")
    print(" "*14 + "[Q1 Q2 Q3 Q4 Q5]")
    for rowNum in range(3, sheet.max_row + 1):
        answer = ""
        if (sheet.cell(row=rowNum, column=4).value == student_for_check[i]):
            # 判斷Q1~Q5的作答情況
            for columnNum in range(6, 11):
                if (sheet.cell(row=rowNum, column=columnNum).value != None):
                    text = sheet.cell(row=rowNum, column=columnNum).value
                    # print(text)
                    text = text.split(" ")
                    if (not text[len(text)-1] == '(correct)'):
                        answer = answer + "0  "
                    else:
                        answer = answer + "1  "
                else:
                    answer = answer + "0  "
            print(" "*15 + answer)


# The rest of the quiz column is assigned 0 score, calculate the final score
# the score equation: take largest 12 score, add up and divided by 0.6
temp = "1"
for i in range(2, quiz_count+1):
    temp = temp + ", " + str(i)

for rowNum in range (2, sheet_score.max_row + 1):
    for columnNum in range(5, 20):
        if (sheet_score.cell(rowNum, columnNum).value == "" or sheet_score.cell(rowNum, columnNum).value == None):
            sheet_score.cell(rowNum, columnNum).value = int(0)
    sheet_score.cell(rowNum, 4).value ="=SUM(LARGE((E" + str(rowNum) + ":S" + str(rowNum) + "), {" + temp + "}))/0.6"

# Save changes to the file
wb_score.save(scoreTable_addr)

# finish reading the slido excel, close to save resource
wb.close()
wb_score.close()
